from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, inspect, select, text
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.main as api
from app.database import (
    Base,
    get_db,
    upgrade_phase1_columns,
    upgrade_phase2_columns,
    upgrade_phase3_columns,
)
from app.models import Candidate, Email
from app.services.extractor import extractor_service


@pytest.fixture
def client(monkeypatch):
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    test_session = sessionmaker(bind=engine, autoflush=False, autocommit=False)

    def override_get_db():
        db = test_session()
        try:
            yield db
        finally:
            db.close()

    api.app.dependency_overrides[get_db] = override_get_db
    monkeypatch.setattr(api, "initialize_database", lambda: None)
    monkeypatch.setattr(api, "deliver_outbox_email", lambda _email_id: None)
    with TestClient(api.app) as test_client:
        test_client.app.state.test_session = test_session
        yield test_client
    api.app.dependency_overrides.clear()
    engine.dispose()


def register_and_login(client, email="owner@example.com", organization="Acme Recruiting"):
    response = client.post(
        "/auth/register",
        json={
            "organization_name": organization,
            "full_name": "Org Owner",
            "email": email,
            "password": "a-long-test-password",
        },
    )
    assert response.status_code == 201
    response = client.post(
        "/auth/token",
        data={"username": email, "password": "a-long-test-password"},
    )
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def test_loopback_vite_origin_is_allowed_by_cors(client):
    response = client.options(
        "/auth/register",
        headers={
            "Origin": "http://127.0.0.1:5173",
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "content-type",
        },
    )

    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == "http://127.0.0.1:5173"


def test_render_frontend_origin_is_allowed_by_cors(client):
    response = client.options(
        "/auth/register",
        headers={
            "Origin": "https://bluepace-ats-frontend.onrender.com",
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "content-type",
        },
    )

    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == "https://bluepace-ats-frontend.onrender.com"


def test_cors_keeps_configured_deployment_origins():
    origins = api.get_allowed_origins("https://bluepace.example.com, https://preview.example.com")

    assert origins == [
        "http://localhost:5173",
        "http://127.0.0.1:5173",
        "https://bluepace-ats-frontend.onrender.com",
        "https://bluepace.example.com",
        "https://preview.example.com",
    ]


def test_health_endpoint(client):
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}
    assert float(response.headers["x-process-time-ms"]) >= 0


def test_fast_resume_validation_does_not_require_auth(client):
    response = client.post(
        "/validate/",
        json={
            "resume_json": {
                "name": "Jane Doe",
                "email": "jane@example.com",
                "skills": ["Python", "SQL"],
                "experience": [],
                "education": [{"degree": "BCA", "university": "Example University"}],
                "is_fresher": True,
            },
            "job_description": "Python SQL developer. Entry-level role.",
        },
    )
    assert response.status_code == 200
    validation = response.json()["validation"]
    assert 0 <= validation["match_score"] <= 100
    assert validation["semantic_mode"] == "lexical_fallback"
    assert float(response.headers["x-process-time-ms"]) >= 0


def test_public_jobs_and_resume_application(client, monkeypatch):
    monkeypatch.setenv("PUBLIC_ORGANIZATION_ID", "1")
    monkeypatch.setattr(api, "_enqueue_resume_ingest", lambda _job_id: SimpleNamespace(id="test-task"))
    headers = register_and_login(client, email="public-owner@example.com", organization="Public ATS")

    created = client.post(
        "/jobs",
        json={
            "title": "Python Developer",
            "description": "Entry-level Python and SQL developer.",
            "status": "open",
            "required_skills": ["Python", "SQL"],
            "minimum_experience_years": 0,
            "fresher_allowed": True,
        },
        headers=headers,
    )
    assert created.status_code == 201
    job_id = created.json()["id"]

    public_jobs = client.get("/public/jobs")
    assert public_jobs.status_code == 200
    assert any(job["id"] == job_id for job in public_jobs.json())

    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("jane@example.com")
    doc.add_paragraph("Python, SQL")
    doc.add_paragraph("BCA | Example University | 2026")
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    applied = client.post(
        f"/public/jobs/{job_id}/apply",
        data={"full_name": "Jane Doe", "email": "jane@example.com"},
        files={"file": ("jane.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert applied.status_code == 200, applied.text
    assert applied.json()["status"] == "accepted"
    assert applied.json()["processing_mode"] in {"external_queue", "db_dispatcher"}

    matches = client.get(f"/jobs/{job_id}/matches", headers=headers)
    assert matches.status_code == 200
    assert len(matches.json()) == 1
    assert 0 <= matches.json()[0]["model_score"] <= 100


def test_job_crud_requires_auth_and_round_trips(client):
    assert client.get("/jobs").status_code == 401
    headers = register_and_login(client)
    payload = {"title": "Senior Engineer", "description": "Build reliable systems"}

    created = client.post("/jobs", json=payload, headers=headers)
    assert created.status_code == 201
    job_id = created.json()["id"]
    assert client.get("/jobs", headers=headers).json()[0]["title"] == "Senior Engineer"
    assert client.get(f"/jobs/{job_id}", headers=headers).status_code == 200

    updated = client.patch(f"/jobs/{job_id}", json={"status": "open"}, headers=headers)
    assert updated.status_code == 200
    assert updated.json()["status"] == "open"
    assert client.delete(f"/jobs/{job_id}", headers=headers).status_code == 204
    assert client.get(f"/jobs/{job_id}", headers=headers).status_code == 404


def test_job_persists_experience_fresher_and_required_skills(client):
    headers = register_and_login(client)
    created = client.post(
        "/jobs",
        json={
            "title": "Backend Engineer",
            "description": "Build APIs",
            "required_skills": ["Python", "PostgreSQL"],
            "minimum_experience_years": 3,
            "fresher_allowed": True,
        },
        headers=headers,
    )

    assert created.status_code == 201
    job_id = created.json()["id"]
    assert created.json()["required_skills"] == ["Python", "PostgreSQL"]
    assert created.json()["minimum_experience_years"] == 3
    assert created.json()["fresher_allowed"] is True

    updated = client.patch(
        f"/jobs/{job_id}",
        json={"required_skills": ["Python"], "minimum_experience_years": 0, "fresher_allowed": False},
        headers=headers,
    )
    assert updated.status_code == 200
    assert updated.json()["required_skills"] == ["Python"]
    assert updated.json()["minimum_experience_years"] == 0
    assert updated.json()["fresher_allowed"] is False


def test_job_analysis_returns_structured_requirements(client, monkeypatch):
    from app.services.matching import matching_service

    monkeypatch.setattr(matching_service, "llm_client", None)
    monkeypatch.setattr(matching_service, "embedding_client", None)
    headers = register_and_login(client)
    job = client.post(
        "/jobs",
        json={
            "title": "Senior Python Engineer",
            "description": "We need a senior Python engineer in Seattle with PostgreSQL experience and a bachelor's degree.",
            "status": "open",
        },
        headers=headers,
    ).json()

    response = client.post(f"/jobs/{job['id']}/analyze", headers=headers)

    assert response.status_code == 200
    analysis = response.json()["jd_analysis"]
    assert "Python" in analysis["required_skills"]
    assert analysis["seniority"] == "senior"
    assert analysis["location"] == "Seattle"
    assert analysis["education"]


def test_explicit_job_requirements_override_jd_inference(client, monkeypatch):
    from app.services.matching import matching_service

    monkeypatch.setattr(matching_service, "llm_client", None)
    monkeypatch.setattr(matching_service, "embedding_client", None)
    headers = register_and_login(client)
    job = client.post(
        "/jobs",
        json={
            "title": "Backend Engineer",
            "description": "Python engineer with 2 years experience based in Seattle.",
            "required_skills": ["Kubernetes"],
            "minimum_experience_years": 7,
            "fresher_allowed": True,
            "status": "open",
        },
        headers=headers,
    ).json()

    response = client.post(f"/jobs/{job['id']}/analyze", headers=headers)

    assert response.status_code == 200
    analysis = response.json()["jd_analysis"]
    assert analysis["required_skills"] == ["Kubernetes"]
    assert analysis["minimum_experience_years"] == 7
    assert analysis["fresher_allowed"] is True


def test_matching_explanations_and_recruiter_override(client, monkeypatch):
    from app.services.matching import matching_service

    monkeypatch.setattr(matching_service, "llm_client", None)
    monkeypatch.setattr(matching_service, "embedding_client", None)
    headers = register_and_login(client)
    job = client.post(
        "/jobs",
        json={
            "title": "Senior Python Engineer",
            "description": "Required: Python and PostgreSQL. 5+ years experience. Bachelor's degree. Based in Seattle.",
            "required_skills": ["Python", "PostgreSQL"],
            "minimum_experience_years": 5,
            "fresher_allowed": True,
            "status": "open",
        },
        headers=headers,
    ).json()
    strong_candidate = client.post(
        "/candidates",
        json={
            "first_name": "Jordan",
            "last_name": "Strong",
            "email": "jordan.strong@example.com",
            "resume_data": {
                "skills": ["Python", "PostgreSQL", "Docker"],
                "location": "Seattle",
                "experience": [{"title": "Backend Engineer", "company": "Acme", "duration": "6 years"}],
                "education": [{"degree": "Bachelor's degree in Computer Science", "university": "State University"}],
            },
        },
        headers=headers,
    ).json()
    weak_candidate = client.post(
        "/candidates",
        json={
            "first_name": "Taylor",
            "last_name": "Weak",
            "email": "taylor.weak@example.com",
            "resume_data": {"skills": ["JavaScript"], "location": "Portland", "experience": [], "education": []},
        },
        headers=headers,
    ).json()

    analyzed = client.post(f"/jobs/{job['id']}/analyze", headers=headers)
    assert analyzed.status_code == 200
    ranked = client.post(f"/jobs/{job['id']}/matches", headers=headers)
    assert ranked.status_code == 200
    matches = ranked.json()
    assert [match["candidate_id"] for match in matches] == [strong_candidate["id"], weak_candidate["id"]]
    assert matches[0]["model_score"] > matches[1]["model_score"]
    assert matches[1]["score_breakdown"]["experience"] == 100
    assert any("Fresher eligibility is allowed" in item for item in matches[1]["explanations"])
    assert {"required_skill_coverage", "preferred_skill_coverage", "experience_alignment", "education_alignment", "location_alignment", "work_mode_alignment", "project_evidence", "semantic_similarity"}.issubset(matches[0]["score_breakdown"])
    assert matches[0]["semantic_mode"] == "lexical_fallback"
    assert {"Python", "PostgreSQL"}.issubset(matches[0]["matched_skills"])
    assert len(matches[0]["cv_summary"]) in {3, 4, 5}
    assert any("Required skills found" in explanation for explanation in matches[0]["explanations"])

    original_score = matches[0]["model_score"]
    feedback = client.patch(
        f"/candidate-matches/{matches[0]['id']}/feedback",
        json={"recruiter_override": 42, "recruiter_note": "Recruiter reviewed portfolio"},
        headers=headers,
    )
    assert feedback.status_code == 200
    assert feedback.json()["model_score"] == original_score
    assert feedback.json()["effective_score"] == 42
    assert feedback.json()["recruiter_note"] == "Recruiter reviewed portfolio"
    persisted = client.get(f"/jobs/{job['id']}/matches", headers=headers).json()
    assert next(match for match in persisted if match["id"] == matches[0]["id"])["effective_score"] == 42

    other_org_headers = register_and_login(client, "other@example.com", "Other Org")
    assert client.get(f"/jobs/{job['id']}/matches", headers=other_org_headers).status_code == 404
    assert client.patch(
        f"/candidate-matches/{matches[0]['id']}/feedback",
        json={"recruiter_override": 90},
        headers=other_org_headers,
    ).status_code == 404


def test_candidate_crud_is_organization_scoped(client):
    first_org_headers = register_and_login(client)
    created = client.post(
        "/candidates",
        json={"first_name": "Casey", "last_name": "Ng", "email": "CASEY@example.com"},
        headers=first_org_headers,
    )
    assert created.status_code == 201
    candidate_id = created.json()["id"]
    assert created.json()["email"] == "casey@example.com"

    second_org_headers = register_and_login(
        client, "other-owner@example.com", "Other Recruiting"
    )
    assert client.get(f"/candidates/{candidate_id}", headers=second_org_headers).status_code == 404
    assert client.get("/candidates", headers=second_org_headers).json() == []

    updated = client.patch(
        f"/candidates/{candidate_id}",
        json={"phone": "+1 555 0100"},
        headers=first_org_headers,
    )
    assert updated.status_code == 200
    assert updated.json()["phone"] == "+1 555 0100"
    assert client.delete(f"/candidates/{candidate_id}", headers=first_org_headers).status_code == 204
    assert any(
        entry["action"] == "candidate.deleted"
        for entry in client.get("/audit-logs?entity_type=candidate", headers=first_org_headers).json()
    )


def test_interviewer_can_read_but_cannot_write(client):
    admin_headers = register_and_login(client)
    created_user = client.post(
        "/users",
        json={
            "full_name": "Interview Panelist",
            "email": "interviewer@example.com",
            "password": "another-test-password",
            "role": "interviewer",
        },
        headers=admin_headers,
    )
    assert created_user.status_code == 201
    token = client.post(
        "/auth/token",
        data={"username": "interviewer@example.com", "password": "another-test-password"},
    ).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    assert client.get("/jobs", headers=headers).status_code == 200
    assert client.post(
        "/jobs",
        json={"title": "Restricted", "description": "Read only"},
        headers=headers,
    ).status_code == 403


def test_duplicate_candidate_email_returns_conflict(client):
    headers = register_and_login(client)
    payload = {"first_name": "Taylor", "last_name": "Lee", "email": "taylor@example.com"}
    assert client.post("/candidates", json=payload, headers=headers).status_code == 201
    assert client.post("/candidates", json=payload, headers=headers).status_code == 409


def test_application_starts_applied_and_can_advance(client):
    headers = register_and_login(client)
    job = client.post(
        "/jobs",
        json={"title": "Product Designer", "description": "Design useful tools", "status": "open"},
        headers=headers,
    ).json()
    candidate = client.post(
        "/candidates",
        json={"first_name": "Morgan", "last_name": "Chen", "email": "morgan@example.com"},
        headers=headers,
    ).json()

    created = client.post(
        "/applications",
        json={"job_id": job["id"], "candidate_id": candidate["id"]},
        headers=headers,
    )
    assert created.status_code == 201
    application_id = created.json()["id"]
    assert created.json()["stage_name"] == "Applied"

    moved = client.post(
        f"/applications/{application_id}/stage",
        json={"stage_name": "Screening"},
        headers=headers,
    )
    assert moved.status_code == 200
    assert moved.json()["stage_name"] == "Screening"


def test_resume_fallback_parses_contact_skills_experience_and_education(monkeypatch):
    monkeypatch.setattr(extractor_service, "client", None, raising=False)
    document = Document()
    for paragraph in [
        "Jane Doe",
        "jane@example.com | +1 415 555 0134",
        "Skills: Python, React, PostgreSQL",
        "Experience",
        "Software Engineer | Acme Labs | 2021 - Present",
        "Built and maintained backend services.",
        "Education",
        "B.S. Computer Science | State University | 2017",
    ]:
        document.add_paragraph(paragraph)
    content = BytesIO()
    document.save(content)

    parsed = extractor_service.extract_to_json(content.getvalue(), "jane.docx")

    assert parsed["email"] == "jane@example.com"
    assert {"Python", "React", "PostgreSQL"}.issubset(parsed["skills"])
    assert parsed["experience"][0]["company"] == "Acme Labs"
    assert parsed["education"][0]["university"] == "State University"


def test_resume_lab_extracts_docx_tables_with_generic_mime(client, monkeypatch):
    monkeypatch.setattr(extractor_service, "client", None, raising=False)
    document = Document()
    table = document.add_table(rows=3, cols=1)
    table.cell(0, 0).text = "Jane Doe"
    table.cell(1, 0).text = "jane@example.com"
    table.cell(2, 0).text = "Skills: Python, React"
    content = BytesIO()
    document.save(content)

    response = client.post(
        "/extract/",
        files={"file": ("jane.docx", content.getvalue(), "application/octet-stream")},
    )

    assert response.status_code == 200
    assert response.json()["data"]["email"] == "jane@example.com"
    assert {"Python", "React"}.issubset(response.json()["data"]["skills"])


def test_resume_lab_rejects_oversized_files(client):
    response = client.post(
        "/extract/",
        files={"file": ("large.pdf", b"x" * (10 * 1024 * 1024 + 1), "application/pdf")},
    )

    assert response.status_code == 413
    assert response.json()["detail"] == "Resume must be 10MB or smaller."


def test_resume_lab_reports_empty_documents_as_unprocessable(client):
    document = Document()
    content = BytesIO()
    document.save(content)

    response = client.post(
        "/extract/",
        files={"file": ("empty.docx", content.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert response.status_code == 422
    assert "No readable text was found" in response.json()["detail"]


def test_pipeline_filters_bulk_actions_csv_and_audit_logs(client):
    headers = register_and_login(client)
    job = client.post(
        "/jobs",
        json={"title": "Platform Engineer", "description": "Build platform", "status": "open"},
        headers=headers,
    ).json()
    candidates = []
    for first_name, email, source, skill in [
        ("Avery", "avery@example.com", "referral", "Python"),
        ("Riley", "riley@example.com", "careers-page", "React"),
    ]:
        response = client.post(
            "/candidates",
            json={
                "first_name": first_name,
                "last_name": "Park",
                "email": email,
                "source": source,
                "resume_data": {"skills": [skill]},
            },
            headers=headers,
        )
        assert response.status_code == 201
        candidates.append(response.json())

    applications = []
    for candidate in candidates:
        response = client.post(
            "/applications",
            json={"job_id": job["id"], "candidate_id": candidate["id"]},
            headers=headers,
        )
        assert response.status_code == 201
        applications.append(response.json())

    filtered = client.get(
        "/applications",
        params={"stage_name": "Applied", "source": "referral", "skill": "Python"},
        headers=headers,
    )
    assert [item["candidate_id"] for item in filtered.json()] == [candidates[0]["id"]]
    applied_date = applications[0]["applied_at"][:10]
    dated = client.get(
        "/applications",
        params={"applied_after": applied_date, "applied_before": applied_date},
        headers=headers,
    )
    assert len(dated.json()) == 2

    moved = client.post(
        "/applications/bulk-stage",
        json={"application_ids": [item["id"] for item in applications], "stage_name": "Interview"},
        headers=headers,
    )
    assert moved.status_code == 200
    assert {item["stage_name"] for item in moved.json()} == {"Interview"}

    exported = client.get("/applications/export.csv?stage_name=Interview", headers=headers)
    assert exported.status_code == 200
    assert "Avery Park" in exported.text and "Riley Park" in exported.text

    rejected = client.post(
        "/applications/bulk-stage",
        json={"application_ids": [applications[0]["id"]], "stage_name": "Rejected"},
        headers=headers,
    )
    assert rejected.status_code == 200
    assert rejected.json()[0]["status"] == "rejected"
    assert client.post(
        f"/applications/{applications[0]['id']}/stage",
        json={"stage_name": "Applied"},
        headers=headers,
    ).status_code == 409

    logs = client.get("/audit-logs?entity_type=application", headers=headers).json()
    assert len(logs) >= 5
    with client.app.state.test_session() as db:
        assert len(db.scalars(select(Email)).all()) >= 5


def test_job_archive_is_audited(client):
    headers = register_and_login(client)
    job = client.post(
        "/jobs",
        json={"title": "Archived Role", "description": "No longer hiring"},
        headers=headers,
    ).json()
    archived = client.post(f"/jobs/{job['id']}/archive", headers=headers)
    assert archived.status_code == 200
    assert archived.json()["status"] == "archived"
    assert any(
        entry["action"] == "job.archived"
        for entry in client.get("/audit-logs?entity_type=job", headers=headers).json()
    )


def test_phase1_schema_upgrade_adds_columns_to_phase0_database():
    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE candidates (id INTEGER PRIMARY KEY)"))
        connection.execute(text("CREATE TABLE applications (id INTEGER PRIMARY KEY)"))
        connection.execute(text("CREATE TABLE emails (id INTEGER PRIMARY KEY)"))

    upgrade_phase1_columns(engine)

    assert "resume_data" in {column["name"] for column in inspect(engine).get_columns("candidates")}
    application_columns = {column["name"]: column for column in inspect(engine).get_columns("applications")}
    assert application_columns["updated_at"]["nullable"] is True
    assert "error_message" in {column["name"] for column in inspect(engine).get_columns("emails")}
    engine.dispose()


def test_phase2_schema_upgrade_adds_jd_and_summary_columns():
    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE jobs (id INTEGER PRIMARY KEY)"))
        connection.execute(text("CREATE TABLE candidates (id INTEGER PRIMARY KEY)"))

    upgrade_phase2_columns(engine)

    job_columns = {column["name"] for column in inspect(engine).get_columns("jobs")}
    candidate_columns = {column["name"] for column in inspect(engine).get_columns("candidates")}
    assert {"jd_analysis", "embedding"}.issubset(job_columns)
    assert "cv_summary" in candidate_columns
    engine.dispose()


def test_phase3_schema_upgrade_adds_job_requirement_columns():
    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE jobs (id INTEGER PRIMARY KEY)"))

    upgrade_phase3_columns(engine)

    columns = {column["name"] for column in inspect(engine).get_columns("jobs")}
    assert {"required_skills", "minimum_experience_years", "fresher_allowed"}.issubset(columns)
    engine.dispose()


def test_embedding_cosine_similarity_drives_semantic_component():
    from types import SimpleNamespace

    from app.services.matching import CandidateMatcher

    matcher = CandidateMatcher()
    job = SimpleNamespace(
        title="Python Engineer",
        description="Build APIs",
        location=None,
        jd_analysis={"required_skills": ["Python"], "preferred_skills": [], "seniority": "unspecified"},
        embedding=[1.0, 0.0],
    )
    candidate = SimpleNamespace(
        first_name="Alex",
        last_name="Vector",
        resume_data={"skills": ["Python"], "experience": [], "education": []},
        embedding=[1.0, 0.0],
        cv_summary=["A", "B", "C"],
    )

    result = matcher.score_candidate(job, candidate)

    assert result["semantic_mode"] == "embedding"
    assert result["score_breakdown"]["semantic"] == 100


def test_candidate_resume_upload_persists_structured_profile(client, monkeypatch, tmp_path):
    monkeypatch.setenv("RESUME_STORAGE_DIR", str(tmp_path))
    monkeypatch.setattr(extractor_service, "client", None, raising=False)
    headers = register_and_login(client)
    candidate = client.post(
        "/candidates",
        json={"first_name": "Jane", "last_name": "Doe", "email": "jane@example.com"},
        headers=headers,
    ).json()
    document = Document()
    for paragraph in [
        "Jane Doe",
        "jane@example.com",
        "Skills: Python, React",
        "Experience",
        "Software Engineer | Acme Labs | 2021 - Present",
        "Education",
        "B.S. Computer Science | State University | 2017",
    ]:
        document.add_paragraph(paragraph)
    content = BytesIO()
    document.save(content)

    response = client.post(
        f"/candidates/{candidate['id']}/resume",
        files={
            "file": (
                "jane.docx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=headers,
    )

    assert response.status_code == 200
    assert response.json()["resume_data"]["experience"][0]["company"] == "Acme Labs"
    assert response.json()["resume_data"]["education"][0]["university"] == "State University"
    assert "raw_text" not in response.json()["resume_data"]
    storage_key = response.json()["resume_storage_key"]
    assert (tmp_path / storage_key).is_file()


def test_csv_export_escapes_spreadsheet_formulas(client):
    headers = register_and_login(client)
    job = client.post(
        "/jobs",
        json={"title": "Data Analyst", "description": "Analyze data", "status": "open"},
        headers=headers,
    ).json()
    candidate = client.post(
        "/candidates",
        json={"first_name": "=SUM(1,1)", "last_name": "Candidate", "email": "formula@example.com"},
        headers=headers,
    ).json()
    assert client.post(
        "/applications",
        json={"job_id": job["id"], "candidate_id": candidate["id"]},
        headers=headers,
    ).status_code == 201

    response = client.get("/applications/export.csv", headers=headers)

    assert response.status_code == 200
    assert "'=SUM(1,1) Candidate" in response.text

def test_advanced_candidate_search_filters_and_boolean_query(client):
    headers = register_and_login(client)
    job = client.post(
        "/jobs",
        json={"title": "Backend Engineer", "description": "Python SQL backend role", "status": "open"},
        headers=headers,
    ).json()

    alice = client.post(
        "/candidates",
        json={
            "first_name": "Alice",
            "last_name": "Python",
            "email": "alice.search@example.com",
            "source": "LinkedIn",
            "resume_data": {
                "skills": ["Python", "SQL"],
                "location": "Hyderabad",
                "notice_period": "30 days",
                "availability": "Immediate",
                "preferred_location": "Hyderabad",
                "work_authorization": "Authorized",
                "highest_education": "B.Tech Computer Science",
                "education": [{"degree": "B.Tech Computer Science", "university": "JNTU"}],
                "experience": [{"title": "Backend Engineer", "company": "Acme Labs", "duration": "4 years"}],
            },
        },
        headers=headers,
    ).json()
    bob = client.post(
        "/candidates",
        json={
            "first_name": "Bob",
            "last_name": "Java",
            "email": "bob.search@example.com",
            "source": "Indeed",
            "resume_data": {
                "skills": ["Python", "SQL", "Java"],
                "location": "Hyderabad",
                "experience": [{"title": "Backend Engineer", "company": "Other Co", "duration": "5 years"}],
            },
        },
        headers=headers,
    ).json()

    with client.app.state.test_session() as db:
        db_candidate = db.get(Candidate, alice["id"])
        db_candidate.tags = ["Immediate", "Backend"]
        db.commit()

    application = client.post(
        "/applications",
        json={"job_id": job["id"], "candidate_id": alice["id"]},
        headers=headers,
    )
    assert application.status_code == 201
    assert client.post(
        f"/applications/{application.json()['id']}/stage",
        json={"stage_name": "Screening"},
        headers=headers,
    ).status_code == 200

    response = client.get(
        "/candidates",
        params={
            "search": "Python AND SQL NOT Java",
            "skill": "Python,SQL",
            "location": "Hyderabad",
            "notice_period": "30 days",
            "education": "Computer Science",
            "job_history": "Acme Labs",
            "availability": "Immediate",
            "preferred_location": "Hyderabad",
            "tags": "Immediate,Backend",
            "stage_name": "Screening",
            "source": "LinkedIn",
            "has_applied_job_id": job["id"],
            "min_experience_years": 3,
            "max_experience_years": 5,
        },
        headers=headers,
    )
    assert response.status_code == 200
    assert [candidate["id"] for candidate in response.json()] == [alice["id"]]

    excluded = client.get(
        "/candidates",
        params={"search": "Python AND SQL NOT Java"},
        headers=headers,
    ).json()
    assert [candidate["id"] for candidate in excluded] == [alice["id"]]


def test_resume_intelligence_extracts_extended_profile_and_non_definitive_signals(monkeypatch):
    monkeypatch.setattr(extractor_service, "client", None, raising=False)
    document = Document()
    for paragraph in [
        "Alex Morgan",
        "alex@example.com | +1 555 123 4567",
        "LinkedIn: linkedin.com/in/alex-morgan | GitHub: github.com/alexmorgan",
        "Current Location: Hyderabad, India",
        "Preferred Location: Bengaluru, India",
        "Notice Period: 30 days",
        "Work Authorization: Authorized to work in India",
        "Skills",
        "Python, React, SQL",
        "Experience",
        "Senior Engineer | Acme Labs | Jan 2021 - Present",
        "Education",
        "M.Tech Computer Science | State University | 2020",
        "Certifications: AWS Certified Cloud Practitioner, Google Cloud Digital Leader",
        "Projects",
        "Recruiting Automation Platform",
    ]:
        document.add_paragraph(paragraph)
    content = BytesIO()
    document.save(content)

    parsed = extractor_service.extract_to_json(content.getvalue(), "alex.docx")

    assert parsed["linkedin"] == "https://linkedin.com/in/alex-morgan"
    assert parsed["github"] == "https://github.com/alexmorgan"
    assert parsed["current_location"] == "Hyderabad, India"
    assert parsed["preferred_location"] == "Bengaluru, India"
    assert parsed["notice_period"] == "30 days"
    assert parsed["work_authorization"] == "Authorized to work in India"
    assert parsed["certifications"]
    assert parsed["companies"] == ["Acme Labs"]
    assert parsed["job_titles"] == ["Senior Engineer"]
    assert parsed["resume_intelligence_version"] == 4
    assert "disclaimer" in parsed["resume_quality"]
    assert "linkedin" not in parsed["resume_quality"]["missing_fields"]
    assert "GitHub" not in parsed["resume_quality"]["optional_missing_fields"]


def test_jd_parser_extracts_preferred_skill_before_sentence_punctuation():
    from app.services.matching import CandidateMatcher

    matcher = CandidateMatcher()
    analysis = matcher.parse_job_description(
        "Python Engineer",
        "Required: Python and SQL. Nice to have: React.",
    )

    assert analysis["required_skills"] == ["Python", "SQL"]
    assert analysis["preferred_skills"] == ["React"]


def test_candidate_ranking_exposes_richer_explainable_factors():
    from types import SimpleNamespace
    from app.services.matching import CandidateMatcher

    matcher = CandidateMatcher()
    job = SimpleNamespace(
        title="Senior Python Engineer",
        description="Build Python APIs with SQL and React.",
        location="Bengaluru, India",
        work_mode="hybrid",
        jd_analysis={
            "required_skills": ["Python", "SQL"],
            "preferred_skills": ["React"],
            "minimum_experience_years": 3,
            "education": "B.Tech Computer Science",
            "location": "Bengaluru, India",
            "work_mode": "hybrid",
            "fresher_allowed": False,
        },
        embedding=None,
    )
    candidate = SimpleNamespace(
        first_name="Alex",
        last_name="Engineer",
        resume_data={
            "skills": ["Python", "SQL", "React"],
            "experience": [{"title": "Backend Engineer", "company": "Acme", "duration": "5 years"}],
            "education": [{"degree": "B.Tech Computer Science", "university": "State University"}],
            "current_location": "Bengaluru, India",
            "preferred_location": "Bengaluru, India",
            "work_mode": "hybrid",
            "projects": [{"name": "API Platform", "description": "Python SQL service"}],
        },
        embedding=None,
        cv_summary=None,
    )

    result = matcher.score_candidate(job, candidate)

    assert result["score_breakdown"]["required_skill_coverage"] == 100
    assert result["score_breakdown"]["preferred_skill_coverage"] == 100
    assert result["score_breakdown"]["experience_alignment"] == 100
    assert result["score_breakdown"]["education_alignment"] > 0
    assert result["score_breakdown"]["location_alignment"] == 100
    assert result["score_breakdown"]["work_mode_alignment"] == 100
    assert result["score_breakdown"]["project_evidence"] == 100
    assert "semantic_similarity" in result["score_breakdown"]
    assert sum(result["score_weights"].values()) == pytest.approx(1.0)
    assert result["matched_required_skills"] == ["Python", "SQL"]
    assert result["matched_preferred_skills"] == ["React"]
    assert result["decision_support_only"] is True


def test_candidate_match_api_returns_explainable_ranking_and_keeps_recruiter_override(client):
    headers = register_and_login(client)
    job = client.post(
        "/jobs",
        json={
            "title": "Python Engineer",
            "description": "Python SQL engineer in Bengaluru. At least 3 years experience. Nice to have: React.",
            "location": "Bengaluru, India",
            "work_mode": "hybrid",
            "status": "open",
            "required_skills": ["Python", "SQL"],
            "minimum_experience_years": 3,
        },
        headers=headers,
    ).json()
    assert job["jd_analysis"]["preferred_skills"] == ["React"]
    candidate = client.post(
        "/candidates",
        json={
            "first_name": "Alex",
            "last_name": "Engineer",
            "email": "alex.ranking@example.com",
            "resume_data": {
                "skills": ["Python", "SQL", "React"],
                "experience": [{"title": "Backend Engineer", "company": "Acme", "duration": "5 years"}],
                "education": [{"degree": "B.Tech Computer Science", "university": "State University"}],
                "current_location": "Bengaluru, India",
                "preferred_location": "Bengaluru, India",
                "work_mode": "hybrid",
                "projects": [{"name": "API Platform", "description": "Python SQL service"}],
            },
        },
        headers=headers,
    ).json()

    ranked = client.post(f"/jobs/{job['id']}/matches", headers=headers)
    assert ranked.status_code == 200
    match = ranked.json()[0]
    assert "required_skill_coverage" in match["score_breakdown"]
    assert "preferred_skill_coverage" in match["score_breakdown"]
    assert "experience_alignment" in match["score_breakdown"]
    assert "project_evidence" in match["score_breakdown"]
    assert "semantic_similarity" in match["score_breakdown"]
    assert set(match["score_weights"]) == {
        "required_skill_coverage", "preferred_skill_coverage", "experience_alignment",
        "education_alignment", "location_alignment", "work_mode_alignment",
        "project_evidence", "semantic_similarity",
    }
    assert match["decision_support_only"] is True
    assert match["matched_required_skills"] == ["Python", "SQL"]
    assert match["matched_preferred_skills"] == ["React"]

    updated = client.patch(
        f"/candidate-matches/{match['id']}/feedback",
        json={"recruiter_override": 42, "recruiter_note": "Recruiter review"},
        headers=headers,
    )
    assert updated.status_code == 200
    assert updated.json()["effective_score"] == 42
    assert updated.json()["recruiter_override"] == 42

    refreshed = client.get(f"/jobs/{job['id']}/matches", headers=headers)
    assert refreshed.status_code == 200
    assert refreshed.json()[0]["effective_score"] == 42


def test_recruiter_assistant_returns_question_specific_structured_answer(client, monkeypatch):
    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
    headers = register_and_login(client)

    job = client.post(
        "/jobs",
        json={
            "title": "Python Backend Engineer",
            "description": "Build APIs with Python and PostgreSQL.",
            "status": "open",
            "required_skills": ["Python", "PostgreSQL"],
            "minimum_experience_years": 2,
        },
        headers=headers,
    ).json()

    candidate = client.post(
        "/candidates",
        json={
            "first_name": "Asha",
            "last_name": "Tester",
            "email": "asha.assistant@example.com",
            "resume_data": {
                "skills": ["Python"],
                "experience": [{"title": "Software Engineer", "company": "Acme", "duration": "3 years"}],
                "education": [{"degree": "B.Tech Computer Science", "university": "Example University"}],
                "projects": ["API monitoring dashboard"],
            },
        },
        headers=headers,
    ).json()

    application = client.post(
        "/applications",
        json={"job_id": job["id"], "candidate_id": candidate["id"]},
        headers=headers,
    ).json()

    response = client.post(
        f"/candidate-tools/applications/{application['id']}/assistant",
        json={"question": "What skills are missing?"},
        headers=headers,
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["intent"] == "missing_skills"
    assert payload["answer"]
    assert payload["key_facts"]
    assert payload["matched_required_skills"] == ["Python"]
    assert payload["missing_required_skills"] == ["PostgreSQL"]
    assert isinstance(payload["interview_questions"], list)
    assert payload["screening_email"]
    assert payload["requirement_explanation"]
    assert payload["guardrails"]
